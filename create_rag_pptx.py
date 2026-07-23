"""Create a 2-slide PowerPoint presentation explaining RAG for beginners."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height

# ── Colour palette ──
INDIGO = RGBColor(0x63, 0x66, 0xF1)
DARK = RGBColor(0x0F, 0x17, 0x2A)
GREY = RGBColor(0x64, 0x74, 0x8B)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
ORANGE = RGBColor(0xF5, 0x7F, 0x17)
SOFT_INDIGO_BG = RGBColor(0xEE, 0xF2, 0xFF)
SOFT_GREEN_BG = RGBColor(0xF0, 0xFD, 0xF4)


def add_accent_bar(slide):
    """Top accent bar."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Emu(68580))
    shape.fill.solid()
    shape.fill.fore_color.rgb = INDIGO
    shape.line.fill.background()


def add_bg(slide):
    """White background."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    font_size=18,
    bold=False,
    color=DARK,
    alignment=PP_ALIGN.LEFT,
    font_name="Calibri",
):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_rounded_rect(
    slide,
    left,
    top,
    width,
    height,
    fill_color,
    border_color=None,
    text="",
    font_size=14,
    font_color=DARK,
    bold=False,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.alignment = PP_ALIGN.LEFT
        tf.margin_left = Emu(91440)
        tf.margin_right = Emu(91440)
        tf.margin_top = Emu(45720)
        tf.margin_bottom = Emu(45720)
    return shape


def add_bullet_card(
    slide, left, top, width, height, emoji, title, lines, fill=LIGHT_BG, border=BORDER
):
    """Card with emoji, title, and bullet lines."""
    shape = add_rounded_rect(slide, left, top, width, height, fill, border)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(137160)
    tf.margin_right = Emu(91440)
    tf.margin_top = Emu(91440)
    tf.margin_bottom = Emu(91440)
    # Clear default paragraph
    for _ in range(len(tf.paragraphs)):
        tf.paragraphs[0].clear()
    # Emoji + Title line
    p = tf.paragraphs[0]
    p.text = f"{emoji}  {title}"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = DARK
    p.font.name = "Calibri"
    p.space_after = Pt(8)
    for line in lines:
        p2 = tf.add_paragraph()
        p2.text = line
        p2.font.size = Pt(14)
        p2.font.color.rgb = GREY
        p2.font.name = "Calibri"
        p2.space_after = Pt(4)
        p2.level = 0
    return shape


# ============================================================
# SLIDE 1 — What Is RAG?
# ============================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide1)
add_accent_bar(slide1)

# Title
add_textbox(
    slide1,
    Inches(0.8),
    Inches(0.5),
    Inches(8),
    Inches(0.5),
    "What Is RAG?",
    font_size=36,
    bold=True,
    color=DARK,
)
add_textbox(
    slide1,
    Inches(0.8),
    Inches(1.0),
    Inches(11),
    Inches(0.6),
    "Retrieval-Augmented Generation — making AI answer questions using YOUR data, not just what it learned from the internet.",
    font_size=18,
    color=GREY,
)

# R A G boxes side by side
letters = [
    ("R", "Retrieval", "Find relevant documents"),
    ("A", "Augment", "Add them as context"),
    ("G", "Generation", "AI answers from context"),
]
for i, (letter, label, desc) in enumerate(letters):
    x = Inches(0.8 + i * 4.1)
    y = Inches(1.7)
    shape = add_rounded_rect(
        slide1,
        x,
        y,
        Inches(3.6),
        Inches(1.2),
        SOFT_INDIGO_BG,
        INDIGO,
        f"  {letter}  —  {label}",
        font_size=22,
        font_color=DARK,
        bold=True,
    )
    tf = shape.text_frame
    p2 = tf.add_paragraph()
    p2.text = f"  {desc}"
    p2.font.size = Pt(15)
    p2.font.color.rgb = GREY
    p2.font.name = "Calibri"
    p2.alignment = PP_ALIGN.LEFT

# Analogy box
add_rounded_rect(
    slide1,
    Inches(0.8),
    Inches(3.2),
    Inches(11.7),
    Inches(1.6),
    SOFT_GREEN_BG,
    GREEN,
    "📁  The Filing Cabinet Analogy",
    font_size=20,
    font_color=DARK,
    bold=True,
)
# Add analogy text as separate textbox
add_textbox(
    slide1,
    Inches(1.0),
    Inches(3.5),
    Inches(11),
    Inches(0.4),
    'Imagine you\'re a personal trainer (the AI) with a filing cabinet of 500 exercise cards (your database). A client asks: "Give me exercises for my calf."',
    font_size=15,
    color=DARK,
)
add_textbox(
    slide1,
    Inches(1.0),
    Inches(3.95),
    Inches(11),
    Inches(0.4),
    "❌  Without RAG: You guess from memory — you might invent exercises or get muscles wrong.",
    font_size=15,
    color=RED,
)
add_textbox(
    slide1,
    Inches(1.0),
    Inches(4.25),
    Inches(11),
    Inches(0.4),
    "✅  With RAG: You pull the 3 most relevant cards and answer using ONLY what's on those cards.",
    font_size=15,
    color=GREEN,
)

# Two bottom cards
add_bullet_card(
    slide1,
    Inches(0.8),
    Inches(5.2),
    Inches(5.6),
    Inches(2.0),
    "🏗️",
    "Phase 1: Indexing (Setup)",
    [
        "Each exercise → 'vector fingerprint' (unique number pattern)",
        "Stored in a vector database (Qdrant)",
        "Happens once, or when new exercises are added",
    ],
)
add_bullet_card(
    slide1,
    Inches(6.9),
    Inches(5.2),
    Inches(5.6),
    Inches(2.0),
    "🔍",
    "Phase 2: Querying (Every Question)",
    [
        "Your question → same fingerprint type",
        "DB finds 3 most similar fingerprints (semantic search)",
        "Gemini answers using ONLY those 3 results",
    ],
)

# Slide number
add_textbox(
    slide1,
    Inches(12.3),
    Inches(7.0),
    Inches(0.8),
    Inches(0.3),
    "1 / 2",
    font_size=12,
    color=GREY,
    alignment=PP_ALIGN.RIGHT,
)


# ============================================================
# SLIDE 2 — Why RAG?
# ============================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2)
add_accent_bar(slide2)

add_textbox(
    slide2,
    Inches(0.8),
    Inches(0.5),
    Inches(8),
    Inches(0.5),
    "Why RAG?",
    font_size=36,
    bold=True,
    color=DARK,
)
add_textbox(
    slide2,
    Inches(0.8),
    Inches(1.0),
    Inches(11),
    Inches(0.5),
    "Three big problems that RAG solves — and what happens without it.",
    font_size=18,
    color=GREY,
)

# 3 benefit cards in a row
benefits = [
    (
        "🚫",
        "No Hallucination",
        [
            "AI is forced to answer only from",
            "retrieved documents.",
            "",
            "Without RAG: AI might invent",
            "fake exercises or wrong muscles.",
        ],
    ),
    (
        "🔄",
        "Always Up-To-Date",
        [
            "Add a new exercise to your JSON",
            "file → re-run vector embedding.",
            "",
            "No retraining, no fine-tuning,",
            "no model updates needed.",
        ],
    ),
    (
        "🧠",
        "Understands Meaning",
        [
            '"calf" matches "Calves", "heel raise",',
            '"gastrocnemius", "lower leg".',
            "",
            "Keyword search would fail.",
            "Semantic search understands intent.",
        ],
    ),
]
for i, (emoji, title, lines) in enumerate(benefits):
    x = Inches(0.8 + i * 4.1)
    add_bullet_card(
        slide2, x, Inches(1.6), Inches(3.7), Inches(2.6), emoji, title, lines
    )

# Without RAG vs With RAG table area
add_textbox(
    slide2,
    Inches(0.8),
    Inches(4.5),
    Inches(5),
    Inches(0.4),
    "Before & After RAG",
    font_size=20,
    bold=True,
    color=DARK,
)

# Without RAG box
add_rounded_rect(
    slide2,
    Inches(0.8),
    Inches(5.0),
    Inches(5.7),
    Inches(2.0),
    RGBColor(0xFE, 0xF2, 0xF2),
    RED,
    "❌  Without RAG (LLM-only)",
    font_size=17,
    font_color=RED,
    bold=True,
)
add_textbox(
    slide2,
    Inches(1.0),
    Inches(5.4),
    Inches(5.3),
    Inches(1.5),
    "• May invent exercises that don't exist\n"
    "• Might recommend wrong equipment\n"
    "• Can't access new data without retraining\n"
    "• No traceability — where did that answer come from?",
    font_size=14,
    color=DARK,
)

# With RAG box
add_rounded_rect(
    slide2,
    Inches(6.9),
    Inches(5.0),
    Inches(5.7),
    Inches(2.0),
    RGBColor(0xF0, 0xFD, 0xF4),
    GREEN,
    "✅  With RAG (Ours)",
    font_size=17,
    font_color=GREEN,
    bold=True,
)
add_textbox(
    slide2,
    Inches(7.1),
    Inches(5.4),
    Inches(5.3),
    Inches(1.5),
    "• Answers grounded in real exercise data\n"
    "• Correct equipment, muscles, body parts\n"
    "• Add exercises anytime — no retraining\n"
    "• Every answer traceable to source documents",
    font_size=14,
    color=DARK,
)

# Slide number
add_textbox(
    slide2,
    Inches(12.3),
    Inches(7.0),
    Inches(0.8),
    Inches(0.3),
    "2 / 2",
    font_size=12,
    color=GREY,
    alignment=PP_ALIGN.RIGHT,
)

# Save
output_path = "RAG_Explained_2_Slides.pptx"
prs.save(output_path)
print(f"✅ Presentation saved to: {output_path}")
